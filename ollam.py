"""Local chat UI for Ollama. Run: python chat_server.py
Env: CHAT_PORT (8000), CHAT_HOST (0.0.0.0 = reachable on your LAN; 127.0.0.1 = this machine only), CHAT_TOKEN (optional password)
Keeps using "local model.py" + chatbot_config.json for persona and saved facts."""
import ast, hashlib, importlib.util, io, json, operator, os, re, secrets, subprocess, sys, tempfile, threading, zipfile
from datetime import datetime
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from urllib.parse import urlparse, parse_qs
import requests

BASE = os.path.dirname(os.path.abspath(__file__))
PORT = int(os.environ.get("CHAT_PORT", "8000"))
HOST = os.environ.get("CHAT_HOST", "0.0.0.0")  # all network interfaces; use 127.0.0.1 for this machine only
TOKEN = os.environ.get("CHAT_TOKEN", "")
OLLAMA = "http://localhost:11434"
VOICE = os.environ.get("PIPER_VOICE", os.path.join(BASE, "en_US-lessac-medium.onnx"))  # Piper voice file

try:  # reuse persona, memory and model name from the terminal bot
    spec = importlib.util.spec_from_file_location("lm", os.path.join(BASE, "local model.py"))
    LM = importlib.util.module_from_spec(spec); spec.loader.exec_module(LM)
    CONFIG = LM.load_config()
except Exception as e:
    print("Persona not loaded:", e); LM = None; CONFIG = {}


def system_message():
    if not LM:
        return None  # no persona: let the model's own Modelfile SYSTEM prompt apply
    facts = LM.load_memory()  # re-read each request so new facts and the time stay fresh
    return {"role": "system", "content": LM.build_system_prompt(CONFIG, datetime.now().astimezone().replace(minute=0, second=0, microsecond=0), facts)}


def speech_text(text):
    """Strip Markdown and code so only readable words are spoken."""
    text = re.sub(r"```.*?```", " ", text, flags=re.S)
    text = re.sub(r"^\s*#+\s*", "", text, flags=re.M)
    text = re.sub(r"\\[()\[\]]|\$\$", "", text)
    text = text.replace("\\times", " times ").replace("\\cdot", " times ").replace("\\approx", " approximately ")
    text = re.sub(r"\^\{?(-?\d+)\}?", r" to the power of \1 ", text)
    text = re.sub(r"[*_`]+", "", text)
    text = re.sub(r"^\s*[-\u2022]\s+", "", text, flags=re.M)
    text = re.sub(r"\s[\u2014\u2013-]\s", ", ", text)
    return re.sub(r"\s+", " ", text).strip()


KOKORO_VOICES = ["af_heart", "af_bella", "af_nicole", "af_sarah", "af_sky", "af_alloy", "af_aoede", "af_jessica", "af_kore", "af_nova", "af_river",
                 "am_adam", "am_michael", "am_echo", "am_eric", "am_fenrir", "am_liam", "am_onyx", "am_puck",
                 "bf_alice", "bf_emma", "bf_isabella", "bf_lily", "bm_daniel", "bm_fable", "bm_george", "bm_lewis"]
KOKORO_PIPES = {}
KOKORO_LOCK = threading.Lock()


def kokoro_ok():
    """True when the kokoro package is installed (checked without loading it)."""
    try:
        return importlib.util.find_spec("kokoro") is not None
    except Exception:
        return False


def kokoro_names():
    return ["kokoro:" + v for v in KOKORO_VOICES] if kokoro_ok() else []


def kokoro_default():
    want = os.environ.get("KOKORO_DEFAULT", "kokoro:af_heart")
    return want if want in kokoro_names() else None


import gc, time

KOKORO_STATE = {"last": 0.0, "watcher": False}
KOKORO_IDLE = float(os.environ.get("KOKORO_IDLE_SECONDS", "300"))  # 0 = never unload


def kokoro_pipe(lang, KPipeline):
    """Return the pipeline for an accent. New accents reuse the first model instead of loading another copy."""
    if lang not in KOKORO_PIPES:
        shared = next((getattr(p, "model", None) for p in KOKORO_PIPES.values() if getattr(p, "model", None) is not None), None)
        try:
            KOKORO_PIPES[lang] = KPipeline(lang_code=lang, model=shared) if shared is not None else KPipeline(lang_code=lang)
        except Exception:
            KOKORO_PIPES[lang] = KPipeline(lang_code=lang)
        if not KOKORO_STATE["watcher"] and KOKORO_IDLE > 0:
            KOKORO_STATE["watcher"] = True
            threading.Thread(target=kokoro_watch, daemon=True).start()
    return KOKORO_PIPES[lang]


def kokoro_unload():
    """Free the Kokoro model. Skipped if a speech job is running right now."""
    if not KOKORO_LOCK.acquire(blocking=False):
        return False
    try:
        if not KOKORO_PIPES:
            return False
        KOKORO_PIPES.clear()
    finally:
        KOKORO_LOCK.release()
    gc.collect()
    try:
        import ctypes
        ctypes.CDLL("libc.so.6").malloc_trim(0)  # hand freed memory back to the system (Linux)
    except Exception:
        pass
    print("Kokoro unloaded to free memory (it reloads on the next Listen).")
    return True


def kokoro_watch():
    while True:
        time.sleep(max(1, min(30, KOKORO_IDLE / 2)))
        if KOKORO_PIPES and time.time() - KOKORO_STATE["last"] > KOKORO_IDLE:
            kokoro_unload()


def kokoro_wav(text, voice, speed):
    """Make a mono 24 kHz WAV with Kokoro. Voice must be one of KOKORO_VOICES (checked by the caller)."""
    import wave
    import numpy as np
    from kokoro import KPipeline
    with KOKORO_LOCK:  # one speech job at a time; the model is not thread-safe
        lang = voice[0]  # a = American English, b = British English
        KOKORO_STATE["last"] = time.time()
        pipe = kokoro_pipe(lang, KPipeline)
        parts = []
        for r in pipe(text, voice=voice, speed=speed):
            audio = getattr(r, "audio", None)
            if audio is None:
                audio = r[2]
            if audio is not None:
                parts.append(audio.detach().cpu().numpy() if hasattr(audio, "detach") else np.asarray(audio))
    KOKORO_STATE["last"] = time.time()
    if not parts:
        raise ValueError("Kokoro made no audio")
    pcm = (np.clip(np.concatenate(parts), -1, 1) * 32767).astype(np.int16)
    buf = io.BytesIO()
    with wave.open(buf, "wb") as w:
        w.setnchannels(1)
        w.setsampwidth(2)
        w.setframerate(24000)
        w.writeframes(pcm.tobytes())
    return buf.getvalue()


def voice_files():
    return sorted(f for f in os.listdir(BASE) if f.endswith(".onnx"))


def fit(msgs, ctx):
    """Drop the oldest turns until the chat fits the context window (about 3.5 characters per token)."""
    budget = int(ctx * 0.8 * 3.5)  # leave room for the reply
    size = lambda ms: sum(len(m["content"]) + 3500 * len(m.get("images", [])) for m in ms)
    while size(msgs) > budget:
        keep = 1 if msgs[0]["role"] == "system" else 0
        if len(msgs) <= keep + 1:
            break
        msgs.pop(keep)
    over = size(msgs) - budget
    if over > 0:  # one huge message left: cut its tail
        last = msgs[-1]
        msgs[-1] = {**last, "content": last["content"][:max(500, len(last["content"]) - over)]}
    return msgs


# ---------- tools and caching (v1.7) ----------
CACHE = os.path.join(BASE, "tts_cache")


def save_cache(path, data):
    """Keep generated speech so pressing Listen twice costs nothing. Oldest files go after 300."""
    os.makedirs(CACHE, exist_ok=True)
    with open(path, "wb") as f:
        f.write(data)
    files = sorted((os.path.join(CACHE, n) for n in os.listdir(CACHE)), key=os.path.getmtime)
    for old in files[:-300]:
        os.remove(old)


def fn(name, desc, props=None, req=()):
    return {"type": "function", "function": {"name": name, "description": desc,
            "parameters": {"type": "object", "properties": props or {}, "required": list(req)}}}


TOOLS = [
    fn("calculator", "Use this for ALL arithmetic, never calculate numbers yourself. Exact. Write power as 4568**12 or 4568^12. Supports + - * / % //, brackets, sqrt, abs, round, min, max, log, ln, log10, sin, cos, tan, factorial, pi, e. One call can hold the whole expression.", {"expression": {"type": "string"}}, ["expression"]),
    fn("current_time", "Get the current local date and time"),
    fn("web_search", "Search the web for current information. Returns titles, links and snippets", {"query": {"type": "string"}}, ["query"]),
]
OPS = {ast.Add: operator.add, ast.Sub: operator.sub, ast.Mult: operator.mul, ast.Div: operator.truediv, ast.Mod: operator.mod,
       ast.FloorDiv: operator.floordiv, ast.Pow: operator.pow, ast.USub: operator.neg, ast.UAdd: operator.pos}


def calc(node):
    if isinstance(node, ast.Expression):
        return calc(node.body)
    if isinstance(node, ast.Constant) and isinstance(node.value, (int, float)):
        return node.value
    if isinstance(node, ast.BinOp) and type(node.op) in OPS:
        if isinstance(node.op, ast.Pow) and abs(calc(node.right)) > 2000:
            raise ValueError("exponent too large")
        return OPS[type(node.op)](calc(node.left), calc(node.right))
    if isinstance(node, ast.UnaryOp) and type(node.op) in OPS:
        return OPS[type(node.op)](calc(node.operand))
    if isinstance(node, ast.Name) and node.id in CONSTS:
        return CONSTS[node.id]
    if isinstance(node, ast.Call) and isinstance(node.func, ast.Name) and node.func.id in FUNCS and not node.keywords:
        return FUNCS[node.func.id](*[calc(a) for a in node.args])
    raise ValueError("only plain arithmetic is allowed")


import math


def _fact(n):
    if n != int(n) or n < 0 or n > 1000:
        raise ValueError("factorial needs a whole number from 0 to 1000")
    return math.factorial(int(n))


CONSTS = {"pi": math.pi, "e": math.e}
FUNCS = {"sqrt": math.sqrt, "abs": abs, "round": round, "min": min, "max": max, "log": math.log, "ln": math.log,
         "log10": math.log10, "sin": math.sin, "cos": math.cos, "tan": math.tan, "factorial": _fact}


def calc_text(expr):
    """Clean up what the model sends, calculate, and add the digit count to big whole numbers."""
    expr = str(expr)
    expr = re.sub(r"(?<=\d),(?=\d{3}(?!\d))", "", expr)
    expr = expr.replace("^", "**").replace("\u00d7", "*").replace("\u00f7", "/").replace("\u2212", "-").strip()
    value = calc(ast.parse(expr, mode="eval"))
    if isinstance(value, float) and value.is_integer() and abs(value) < 1e15:
        value = int(value)
    text = str(value)
    if isinstance(value, int) and len(str(abs(value))) > 12:
        text += f"  ({len(str(abs(value)))} digits)"
    return text


def web_search(query):
    from html import unescape
    from urllib.parse import unquote
    page = requests.get("https://lite.duckduckgo.com/lite/", params={"q": query}, headers={"User-Agent": "Mozilla/5.0"}, timeout=15).text
    items = re.findall(r"<a[^>]*href=['\"]([^'\"]+)['\"][^>]*class=['\"]result-link['\"][^>]*>(.*?)</a>", page, re.S)
    snips = re.findall(r"class=['\"]result-snippet['\"][^>]*>(.*?)</td>", page, re.S)
    out = []
    for i, (href, title) in enumerate(items[:5]):
        hit = re.search(r"uddg=([^&]+)", href)
        snippet = unescape(re.sub(r"<[^>]+>|\s+", " ", snips[i])).strip() if i < len(snips) else ""
        out.append(f"{i + 1}. {unescape(re.sub(r'<[^>]+>', '', title)).strip()}\n{unquote(hit.group(1)) if hit else href}\n{snippet}")
    return "\n\n".join(out) or "No results."


TOOL_HINT = {"calculator": "\n(Write this number out in full as plain text. Do not use LaTeX or scientific notation.)"}


def tool_reply(name, args):
    out = run_tool(name, args)
    return out if out.startswith("Tool error") else out + TOOL_HINT.get(name, "")


def run_tool(name, args):
    try:
        if name == "calculator":
            return calc_text(args.get("expression", ""))
        if name == "current_time":
            return datetime.now().astimezone().strftime("%A, %B %d, %Y, %H:%M %Z")
        if name == "web_search":
            return web_search(str(args.get("query", "")))[:3000]
        return f"Unknown tool: {name}"
    except Exception as error:
        return f"Tool error: {error}"


# ---------- effort routing (v1.8) ----------
THINK_CAPS = {}
HARD = re.compile(r"\b(why|how|explain|compare|debug|fix|error|bug|plan|prove|calculate|solve|analy[sz]e|code|script|function|essay|summari[sz]e|difference|step)\b", re.I)


def can_think(model):
    """Ask Ollama once per model whether it supports thinking (qwen2.5 does not)."""
    if model not in THINK_CAPS:
        try:
            info = requests.post(OLLAMA + "/api/show", json={"model": model}, timeout=5).json()
            THINK_CAPS[model] = "thinking" in info.get("capabilities", [])
        except Exception:
            THINK_CAPS[model] = False
    return THINK_CAPS[model]


SEE_CAPS = {}


def can_see(model):
    """Ask Ollama once per model whether it can look at pictures (qwen2.5 cannot)."""
    if model not in SEE_CAPS:
        try:
            info = requests.post(OLLAMA + "/api/show", json={"model": model}, timeout=5).json()
            SEE_CAPS[model] = "vision" in info.get("capabilities", [])
        except Exception:
            SEE_CAPS[model] = False
    return SEE_CAPS[model]


def decide_think(req, msgs, model):
    """Fast = never, Deep = always, Auto = only for long, code-like or how/why style questions."""
    effort = req.get("effort", "auto")
    if effort == "fast" or not can_think(model):
        return False
    if effort == "deep":
        return True
    last = next((m["content"] for m in reversed(msgs) if m["role"] == "user"), "")
    return len(last) > 400 or "```" in last or (len(last) > 40 and bool(HARD.search(last)))


# ---------- GPU routing (v1.9) ----------
CPU_ONLY = [w.strip().lower() for w in os.environ.get("CPU_ONLY_MODELS", "9b,heretic-writer").split(",") if w.strip() and w.strip().lower() != "none"]


def gpu_options(model):
    """Models too big for the iGPU run on the CPU (num_gpu 0); others are left to Ollama's own placement."""
    name = (model or "").lower()
    return {"num_gpu": 0} if any(w in name for w in CPU_ONLY) else {}


# ---------- export helpers (v1.3) ----------
MAX_EXPORT = 8 * 1024 * 1024
LANG_EXT = {"python": "py", "py": "py", "javascript": "js", "js": "js", "jsx": "jsx", "typescript": "ts", "ts": "ts",
            "tsx": "tsx", "bash": "sh", "sh": "sh", "shell": "sh", "zsh": "sh", "powershell": "ps1", "bat": "bat",
            "html": "html", "css": "css", "scss": "scss", "json": "json", "yaml": "yml", "yml": "yml", "toml": "toml",
            "ini": "ini", "xml": "xml", "sql": "sql", "java": "java", "kotlin": "kt", "c": "c", "cpp": "cpp",
            "c++": "cpp", "csharp": "cs", "cs": "cs", "go": "go", "rust": "rs", "php": "php", "ruby": "rb",
            "swift": "swift", "dart": "dart", "lua": "lua", "markdown": "md", "md": "md", "text": "txt",
            "txt": "txt", "csv": "csv"}
KNOWN_EXT = set(LANG_EXT.values()) | {"cfg", "conf", "env", "gradle", "vue", "svelte", "svg"}
CTRL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
FENCE = re.compile(r"```([\w+#.-]*)[^\n]*\n(.*?)(?:```|\Z)", re.S)
NAME_LINE = re.compile(r"^\s*(?:#|//|--|;|<!--|/\*)\s*(?:file(?:name)?\s*[:=]\s*)?([\w./\\-]+\.[A-Za-z0-9]{1,6})\s*(?:-->|\*/)?\s*$", re.I)
NAME_TICKS = re.compile(r"[`*\"]([\w./\\-]+\.[A-Za-z0-9]{1,6})[`*\"]")


def clean_messages(raw):
    """Keep only well-formed user/assistant text from the browser."""
    out = []
    for m in (raw if isinstance(raw, list) else [])[-500:]:
        if isinstance(m, dict) and m.get("role") in ("user", "assistant") and isinstance(m.get("content"), str):
            text = CTRL.sub("", m["content"])[:200000]
            if text.strip():
                out.append({"role": m["role"], "content": text})
    return out


def safe_name(name):
    """Flat, harmless file name: no folders, no leading dots, only simple characters."""
    name = os.path.basename(name.replace("\\", "/"))
    return re.sub(r"[^\w.\-]", "_", name).lstrip(".")


def code_blocks(messages):
    """Return [(filename, text)] for every fenced code block in the assistant replies."""
    files, used = [], set()
    for msg in messages:
        if msg["role"] != "assistant":
            continue
        text = msg["content"]
        for m in FENCE.finditer(text):
            lang, body = m.group(1).lower(), m.group(2)
            if not body.strip():
                continue
            hit = NAME_LINE.match(body.split("\n", 1)[0])  # a first-line comment such as "# server.py"
            name = hit.group(1) if hit else None
            if not name:  # otherwise a `file.ext` mentioned on the line just before the block
                ticks = NAME_TICKS.findall(text[:m.start()].rstrip().rsplit("\n", 1)[-1])
                name = ticks[-1] if ticks else None
            name = safe_name(name) if name else ""
            if name.rsplit(".", 1)[-1].lower() not in KNOWN_EXT:
                name = f"snippet_{len(files) + 1}.{LANG_EXT.get(lang, 'txt')}"
            stem, ext = os.path.splitext(name)
            n = 2
            while name.lower() in used:
                name, n = f"{stem}_{n}{ext}", n + 1
            used.add(name.lower())
            files.append((name, body.rstrip("\n") + "\n"))
    return files[:100]


def to_markdown(messages, title):
    parts = [f"# {title}", f"_Exported {datetime.now():%Y-%m-%d %H:%M}_"]
    for m in messages:
        parts += ["## " + ("You" if m["role"] == "user" else "Assistant"), m["content"].strip()]
    return "\n\n".join(parts) + "\n"


def make_zip(entries):
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for name, data in entries:
            z.writestr(name, data)
    return buf.getvalue()


def build_docx(messages, title):
    """Markdown-ish chat to Word. Needs: pip install python-docx"""
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt

    doc = Document()
    doc.add_heading(title, 0)

    def inline(par, text):
        for piece in re.split(r"(\*\*[^*\n]+\*\*|`[^`\n]+`)", text):
            if piece.startswith("**") and piece.endswith("**") and len(piece) > 4:
                par.add_run(piece[2:-2]).bold = True
            elif piece.startswith("`") and piece.endswith("`") and len(piece) > 2:
                par.add_run(piece[1:-1]).font.name = "Consolas"
            elif piece:
                par.add_run(piece)

    def code(text):
        par = doc.add_paragraph()
        shd = OxmlElement("w:shd")
        for k, v in (("val", "clear"), ("color", "auto"), ("fill", "F2F4F8")):
            shd.set(qn("w:" + k), v)
        par._p.get_or_add_pPr().append(shd)
        run = par.add_run(text.rstrip("\n"))
        run.font.name, run.font.size = "Consolas", Pt(9)

    def table(rows):
        cols = max(len(r) for r in rows)
        t = doc.add_table(rows=len(rows), cols=cols)
        t.style = "Table Grid"
        for i, r in enumerate(rows):
            for j in range(cols):
                par = t.cell(i, j).paragraphs[0]
                inline(par, r[j] if j < len(r) else "")
                if i == 0:
                    for run in par.runs:
                        run.bold = True

    for msg in messages:
        doc.add_heading("You" if msg["role"] == "user" else "Assistant", 2)
        lines, i = msg["content"].split("\n"), 0
        while i < len(lines):
            line = lines[i]
            if line.strip().startswith("```"):
                j, block = i + 1, []
                while j < len(lines) and not lines[j].strip().startswith("```"):
                    block.append(lines[j])
                    j += 1
                code("\n".join(block))
                i = j + 1
                continue
            if line.lstrip().startswith("|"):
                rows = []
                while i < len(lines) and lines[i].lstrip().startswith("|"):
                    cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                    if not all(re.fullmatch(r":?-+:?", c) for c in cells):
                        rows.append(cells)
                    i += 1
                if rows:
                    table(rows)
                continue
            h = re.match(r"(#{1,6})\s+(.*)", line)
            b = re.match(r"\s*[-*\u2022]\s+(.*)", line)
            n = re.match(r"\s*(\d+[.)])\s+(.*)", line)
            if re.fullmatch(r"\s*(-{3,}|\*{3,})\s*", line):
                pass
            elif h:
                doc.add_heading(re.sub(r"[*`]", "", h.group(2)), min(len(h.group(1)) + 2, 4))
            elif b:
                inline(doc.add_paragraph(style="List Bullet"), b.group(1))
            elif n:
                par = doc.add_paragraph()
                par.paragraph_format.left_indent = Pt(18)
                inline(par, f"{n.group(1)} {n.group(2)}")
            elif line.strip():
                inline(doc.add_paragraph(), line.strip())
            i += 1
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


PAGE = r'''<!doctype html><html lang="en"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1,viewport-fit=cover,interactive-widget=resizes-content"><meta name="theme-color" content="#102343"><meta name="mobile-web-app-capable" content="yes"><meta name="apple-mobile-web-app-capable" content="yes"><title>Local assistant</title>
<style>
:root{--navy:#102343;--blue:#2f4d80;--amber:#fdbd27;--page:#f5f7fb;--muted:#7890b2;--line:#e4eaf3;--soft:#fff6d9;--side:#102343;--text:#102343;--bub:#102343;--bubt:#fff;--code:#102343;--codet:#e8eefb;--cbb:transparent;--card:#fff;--icode:#e6ecf6;--thinkt:#5c5230;--thinkh:#6b5410;--link:#2f4d80;--shadow:#10234312}
:root[data-theme=dark]{color-scheme:dark;--page:#0c1424;--side:#08101f;--text:#e3eaf6;--bub:#2f4d80;--code:#060b16;--cbb:#25365a;--card:#142039;--icode:#1f2d4a;--line:#25365a;--muted:#8aa0c4;--soft:#2a2410;--thinkt:#d9cc9a;--thinkh:#f3d77a;--link:#8fb0e8;--shadow:#0007}
*{box-sizing:border-box}html,body{height:100%;margin:0}
body{font:16px/1.65 system-ui,-apple-system,"Segoe UI",sans-serif;color:var(--text);background:var(--page);display:grid;grid-template-columns:260px 1fr}
aside{background:var(--side);color:#dbe5f5;padding:22px 18px;display:flex;flex-direction:column;gap:18px}
.brand{display:flex;gap:11px;align-items:center;font-weight:650;font-size:17px;color:#fff}
.mark{width:34px;height:34px;border-radius:10px;background:var(--amber);color:var(--navy);display:grid;place-items:center;font-weight:800}
button,select,textarea{font:inherit}
.new{background:var(--amber);color:var(--navy);border:0;border-radius:10px;padding:11px;font-weight:650;cursor:pointer}
label.f{display:block;font-size:13px;color:#9fb3d6;margin-bottom:6px}
select{width:100%;padding:10px;border-radius:9px;border:1px solid #2f4d80;background:#17305a;color:#fff}
.sw{display:flex;gap:10px;align-items:flex-start;cursor:pointer}.sw input{position:absolute;opacity:0}
.sw i{flex:none;width:38px;height:22px;border-radius:12px;background:#2f4d80;position:relative;transition:.15s}
.sw i:after{content:"";position:absolute;top:3px;left:3px;width:16px;height:16px;border-radius:50%;background:#fff;transition:.15s}
.sw input:checked+i{background:var(--amber)}.sw input:checked+i:after{left:19px;background:var(--navy)}
.sw input:focus-visible+i,select:focus-visible,button:focus-visible,textarea:focus-visible{outline:2px solid var(--amber);outline-offset:2px}
.sw small{display:block;color:#9fb3d6;font-size:12.5px;line-height:1.4}
.foot{margin-top:auto;font-size:12.5px;color:#9fb3d6}
main{display:flex;flex-direction:column;height:100vh;min-width:0}
#log{flex:1;overflow:auto;padding:36px 20px 10px}
.wrap{max-width:760px;margin:0 auto}
.empty{margin:18vh auto 0;max-width:760px;text-align:center}.empty h1{font-size:30px;margin:0 0 8px}.empty p{color:var(--muted);margin:0}
.m{max-width:760px;margin:0 auto 22px}
.user .bub{margin-left:auto;width:fit-content;max-width:88%;background:var(--bub);color:var(--bubt);padding:11px 16px;border-radius:18px 18px 4px 18px;white-space:pre-wrap;overflow-wrap:anywhere}
.body{white-space:pre-wrap;overflow-wrap:anywhere}
.cb{background:var(--code);border:1px solid var(--cbb);border-radius:12px;margin:10px 0;overflow:hidden}
.cbh{display:flex;justify-content:space-between;align-items:center;padding:7px 10px 0 16px;font:12px ui-monospace,Menlo,Consolas,monospace;color:#9fb3d6}
.cbh button{background:none;border:1px solid #33496f;color:#c9d6ee;border-radius:7px;padding:1px 10px;font-size:12px;cursor:pointer}.cbh button:hover{background:#ffffff14}
.body pre{white-space:pre;overflow:auto;margin:0;padding:6px 16px 14px;color:var(--codet);font:13.5px/1.55 ui-monospace,Menlo,Consolas,monospace}
.body code{background:var(--icode);padding:1px 6px;border-radius:5px;font:.9em ui-monospace,Menlo,Consolas,monospace}.body pre code{background:none;padding:0}
.meta{font-size:12.5px;color:var(--muted);margin-top:4px}
.think{background:var(--soft);border-left:3px solid var(--amber);border-radius:0 10px 10px 0;margin-bottom:12px;font-size:14px}
.think summary{cursor:pointer;padding:9px 14px;color:var(--thinkh);font-weight:600;list-style:none}
.think summary::before{content:"";display:inline-block;width:8px;height:8px;border-radius:50%;background:var(--amber);margin-right:9px}
.think.live summary::before{animation:p 1s ease-in-out infinite}
@keyframes p{50%{opacity:.25;transform:scale(.7)}}
.think div{padding:0 14px 12px;color:var(--thinkt);white-space:pre-wrap;max-height:240px;overflow:auto}
.dock{padding:8px 20px 18px}
.box{max-width:760px;margin:0 auto;background:var(--card);border:1px solid var(--line);border-radius:18px;padding:10px;display:flex;gap:8px;align-items:flex-end;box-shadow:0 6px 24px var(--shadow)}
.box:focus-within{border-color:var(--link)}
textarea{flex:1;border:0;outline:0;resize:none;max-height:200px;padding:8px 10px;color:var(--text);background:none}
.ghost{background:none;border:1px solid var(--line);color:var(--link);border-radius:10px;padding:8px 12px;cursor:pointer}
.go{background:var(--amber);color:var(--navy);border:0;border-radius:10px;padding:9px 18px;font-weight:700;cursor:pointer}.go.stop{background:var(--bub);color:var(--bubt)}
.lis{padding:3px 11px;font-size:12.5px}.acts{display:flex;gap:6px;flex-wrap:wrap;margin-top:6px}
.eg{display:grid;grid-template-columns:1fr 1fr;gap:6px}.exp .ghost{padding:6px 4px;font-size:13px;color:#dbe5f5;border-color:#2f4d80}.exp .ghost:hover{background:#17305a}
main{position:relative}
#chats{max-height:30vh;overflow:auto;display:flex;flex-direction:column;gap:4px}
.ci{display:flex;justify-content:space-between;gap:8px;padding:7px 10px;border-radius:8px;cursor:pointer;font-size:13.5px;color:#c9d6ee}
.ci.on{background:#17305a;color:#fff}.ci span{overflow:hidden;text-overflow:ellipsis;white-space:nowrap}.ci b{opacity:.55}.ci b:hover{opacity:1}
.vc{display:flex;gap:6px}.vc select{padding:6px;font-size:12.5px}.vc input{width:84px}
#jump{position:absolute;right:28px;bottom:118px;background:var(--card);box-shadow:0 4px 16px var(--shadow)}
.body a{color:var(--link)}
@media(max-width:760px){#chats,.vc{display:none}}
.hint{max-width:760px;margin:6px auto 0;font-size:12px;color:var(--muted);text-align:center}
@media(max-width:760px){body{grid-template-columns:1fr;grid-template-rows:auto 1fr}aside{flex-direction:row;flex-wrap:wrap;align-items:center;padding:12px}.foot,.sw small,label.f,.brand span{display:none}main{height:auto;min-height:0}#log{padding-top:20px}.empty{margin-top:8vh}}
#bar,#scrim{display:none}
@media(max-width:760px){
body{display:block;height:100dvh;overflow:hidden}
aside{position:fixed;top:0;bottom:0;left:0;width:min(86vw,320px);flex-direction:column;flex-wrap:nowrap;align-items:stretch;overflow-y:auto;padding:calc(18px + env(safe-area-inset-top)) 16px calc(18px + env(safe-area-inset-bottom));transform:translateX(-103%);transition:transform .22s ease;z-index:30;box-shadow:6px 0 30px #0006}
body.menu aside{transform:none}
#scrim{position:fixed;inset:0;background:#0009;z-index:25}body.menu #scrim{display:block}
aside .foot,aside .sw small,aside label.f{display:block}aside .brand span{display:inline}
#chats{display:flex;max-height:34vh}.vc{display:flex}.vc input{flex:1}
.ci{padding:11px 12px;font-size:15px}
main{height:100dvh}
#bar{display:flex;align-items:center;gap:10px;flex:none;padding:calc(8px + env(safe-area-inset-top)) 12px 8px;background:var(--side);color:#fff}
#ttl{flex:1;min-width:0;font-weight:650;overflow:hidden;text-overflow:ellipsis;white-space:nowrap}
#bar .ghost{color:#fff;border-color:#33496f;min-width:44px;min-height:42px;font-size:19px;line-height:1;padding:0 10px}
#log{padding:14px 12px 6px;overscroll-behavior:contain}
.m{margin-bottom:18px}.user .bub{max-width:94%}
.dock{padding:6px 10px calc(8px + env(safe-area-inset-bottom))}
.box{border-radius:16px;padding:8px}textarea{font-size:16px}
.ghost,.go{min-height:42px}.lis{min-height:36px;padding:4px 12px}.acts{gap:8px}
.cb pre{font-size:12.5px}.cbh button{min-height:30px}
.hint{font-size:11px}
#jump{bottom:calc(96px + env(safe-area-inset-bottom));right:14px}
.empty{margin-top:6vh}.empty h1{font-size:24px}
}
@media(prefers-reduced-motion:reduce){*{animation:none!important;transition:none!important}}
</style><script>document.documentElement.dataset.theme=localStorage.theme||(matchMedia('(prefers-color-scheme:dark)').matches?'dark':'light')</script></head><body>
<div id="scrim"></div><aside><div class="brand"><div class="mark">A</div><span>Local assistant</span></div>
<button class="new" id="new">New chat</button><div id="chats"></div>
<div><label class="f" for="model">Model</label><select id="model"></select></div>
<div><label class="f" for="ctx">Memory (context)</label><select id="ctx"><option value="4096">4k - fastest</option><option value="8192">8k</option><option value="16384">16k</option><option value="32768">32k - slow, heavy</option></select></div>
<div><label class="f" for="think">Thinking</label><select id="think"><option value="auto">Auto - only for hard questions</option><option value="fast">Fast - never</option><option value="deep">Deep - always</option></select></div>
<label class="sw"><input type="checkbox" id="speak"><i></i><span>Speak replies<small>Reads each answer aloud with Kokoro or Piper.</small></span></label>
<label class="sw"><input type="checkbox" id="tools"><i></i><span>Use tools<small>Calculator, clock, web search.</small></span></label>
<label class="sw"><input type="checkbox" id="vis"><i></i><span>Model sees pictures<small>Off = read only the text in a picture (OCR).</small></span></label>
<div class="vc"><select id="vsel" title="Voice"></select><input type="range" id="vspd" min="0.7" max="1.5" step="0.05" value="1" title="Speech speed"></div>
<label class="sw"><input type="checkbox" id="dark"><i></i><span>Dark mode<small>Easier on the eyes at night.</small></span></label>
<div class="exp"><label class="f">Export chat</label><div class="eg"><button class="ghost" data-x="md">Markdown</button><button class="ghost" data-x="docx">Word</button><button class="ghost" data-x="code">Code zip</button><button class="ghost" data-x="bundle">Everything</button></div></div>
<div class="foot">Runs on your own machine.</div></aside>
<main><header id="bar"><button class="ghost" id="menu" aria-label="Open menu">☰</button><div id="ttl">Local assistant</div><button class="ghost" id="nw" aria-label="New chat">+</button></header><div id="log"></div><button class="ghost" id="jump" hidden>Latest</button>
<div class="dock"><div class="box"><input type="file" id="file" hidden accept="image/*,.txt,.md,.csv,.json,.py,.js,.html,.xml,.log">
<button class="ghost" id="attach" title="Attach a text file">Attach</button>
<textarea id="in" rows="1" placeholder="Message the assistant" autofocus></textarea>
<button class="go" id="go">Send</button></div>
<div class="hint" id="hint">Enter to send, Shift+Enter for a new line</div></div></main>
<script>
const $=id=>document.getElementById(id),log=$('log'),ta=$('in'),go=$('go'),sel=$('model'),thk=$('think');
let chats=JSON.parse(localStorage.chats||'null')||[{id:Date.now(),title:'New chat',msgs:JSON.parse(localStorage.chat||'[]')}],cur=+localStorage.cur,att=null,ctrl=null;
if(!chats.some(c=>c.id===cur))cur=chats[0].id;
let chat=chats.find(c=>c.id===cur).msgs;
thk.value=localStorage.effort||'auto';thk.onchange=()=>localStorage.effort=thk.value;
const spkChk=$('speak');spkChk.checked=localStorage.speak==='1';spkChk.onchange=()=>{localStorage.speak=spkChk.checked?'1':'0';if(!spkChk.checked)hush()};
const esc=s=>s.replace(/[&<>]/g,c=>({'&':'&amp;','<':'&lt;','>':'&gt;'}[c]));
const SUP={'0':'\u2070','1':'\u00b9','2':'\u00b2','3':'\u00b3','4':'\u2074','5':'\u2075','6':'\u2076','7':'\u2077','8':'\u2078','9':'\u2079','-':'\u207b'};
const tex=s=>s.split('```').map((p,i)=>i%2?p:p.replace(/\\left|\\right/g,'').replace(/\\\(|\\\)|\\\[|\\\]|\$\$/g,'').replace(/\\text\{([^}]*)\}/g,'$1').replace(/\\frac\{([^}]*)\}\{([^}]*)\}/g,'($1)/($2)').replace(/\\sqrt\{([^}]*)\}/g,'\u221a($1)').replace(/\\times/g,'\u00d7').replace(/\\cdot/g,'\u00b7').replace(/\\div/g,'\u00f7').replace(/\\approx/g,'\u2248').replace(/\\leq?/g,'\u2264').replace(/\\geq?/g,'\u2265').replace(/\\neq/g,'\u2260').replace(/\\pm/g,'\u00b1').replace(/\\pi/g,'\u03c0').replace(/\^\{(-?\d+)\}|(?<=[\w)])\^(-?\d+)/g,(_,a,b)=>[...(a||b)].map(c=>SUP[c]).join(''))).join('```');
const md=s=>{const k=[];let t=esc(tex(s)).replace(/\n*```([\w+#.-]*)[^\n]*\n([\s\S]*?)(```|$)\n*/g,(_,l,c)=>{k.push(`<div class="cb"><div class="cbh"><span>${l||'code'}</span><button class="cp">Copy</button></div><pre><code>${c.replace(/\n$/,'')}</code></pre></div>`);return`\u0000${k.length-1}\u0000`});
 t=t.replace(/^(\s*)[-*] /gm,'$1\u2022 ').replace(/`([^`\n]+)`/g,'<code>$1</code>').replace(/\*\*([^*\n]+)\*\*/g,'<b>$1</b>').replace(/(^|[^*\w])\*([^*\n]+)\*/g,'$1<i>$2</i>').replace(/\[([^\]\n]+)\]\((https?:[^)\s]+)\)/g,'<a href="$2" target="_blank" rel="noopener">$1</a>').replace(/^#{1,3} (.+)$/gm,'<b>$1</b>');
 return t.replace(/\u0000(\d+)\u0000/g,(_,i)=>k[i])};
function paint(d,m,live){const t=d.querySelector('.think');
 if(m.think){t.hidden=false;t.querySelector('div').textContent=m.think;const going=live&&!m.content;
  t.querySelector('summary').textContent=going?'Thinking...':`Thought for ${m.secs||0}s`;t.classList.toggle('live',going);
  if(going)t.open=true;else if(!m.shut){t.open=false;m.shut=1}}
 d.querySelector('.body').innerHTML=md(m.content);d.querySelector('.meta').textContent=(m.err?m.err+' \u00b7 ':'')+(m.stats||'');const z=d.querySelector('[data-a=z]');if(z)z.hidden=!m.content.includes('```');const k=d.querySelector('[data-a=k]');if(k)k.hidden=!(m.err||m.stats==='Stopped')}
function addMsg(m,live){const d=document.createElement('div');d.className='m '+m.role;
 if(m.role==='user'){d.innerHTML='<div class="bub"></div>';d.firstChild.textContent=m.shown??m.content}
 else{d.innerHTML='<details class="think" hidden><summary></summary><div></div></details><div class="body"></div><div class="meta"></div><div class="acts"><button class="ghost lis" data-a="l">Listen</button><button class="ghost lis" data-a="c">Copy reply</button><button class="ghost lis" data-a="z" hidden>Code zip</button><button class="ghost lis" data-a="r">Regenerate</button><button class="ghost lis" data-a="k" hidden>Continue</button></div>';paint(d,m,live);d.querySelector('[data-a=r]').onclick=regen;d.querySelector('[data-a=k]').onclick=()=>{ta.value='Continue exactly where you stopped.';send()};
  d.querySelector('[data-a=l]').onclick=()=>speak(m.content);
  d.querySelector('[data-a=c]').onclick=async e=>{const b=e.target;if(await copyText(m.content)){b.textContent='Copied';setTimeout(()=>{b.textContent='Copy reply'},1200)}};
  d.querySelector('[data-a=z]').onclick=()=>dl('code',[{role:'assistant',content:m.content}])}
 log.appendChild(d);log.scrollTop=log.scrollHeight;return d}
function render(){log.innerHTML='';if(!chat.length)log.innerHTML='<div class="empty"><h1>What do you want to work on?</h1><p>Ask a question, paste code, or attach a text file.</p></div>';chat.forEach(m=>addMsg(m,false));tidy();list()}
const save=()=>{const c=chats.find(x=>x.id===cur);c.msgs=chat.slice(-80);chat=c.msgs;
 if(c.title==='New chat'){const f=chat.find(m=>m.role==='user');if(f)c.title=(f.shown||f.content).slice(0,40)}
 try{localStorage.chats=JSON.stringify(chats,(k,v)=>k==='images'?undefined:v);localStorage.cur=cur}catch(_){}list()};
function list(){setTitle();const box=$('chats');box.innerHTML='';chats.forEach(c=>{const d=document.createElement('div');d.className='ci'+(c.id===cur?' on':'');
 const s=document.createElement('span');s.textContent=c.title;s.title='Double-click to rename';const x=document.createElement('b');x.textContent='\u00d7';x.title='Delete chat';
 d.append(s,x);d.onclick=()=>pick(c.id);
 s.ondblclick=e=>{e.stopPropagation();const t=prompt('Rename chat',c.title);if(t){c.title=t.slice(0,60);save()}};
 x.onclick=e=>{e.stopPropagation();if(!confirm('Delete "'+c.title+'"?'))return;chats=chats.filter(k=>k.id!==c.id);
  if(!chats.length)chats=[{id:Date.now(),title:'New chat',msgs:[]}];if(c.id===cur)pick(chats[0].id);save()};box.appendChild(d)})}
function pick(id){document.body.classList.remove('menu');hush();if(ctrl)ctrl.abort();cur=id;chat=chats.find(c=>c.id===id).msgs;localStorage.cur=cur;render()}
function newChat(){const e=chats.find(c=>c.id===cur);if(e&&!e.msgs.length){ta.focus();return}chats.unshift({id:Date.now(),title:'New chat',msgs:[]});pick(chats[0].id);save()}
function tidy(){const all=log.querySelectorAll('.m.assistant');all.forEach((d,i)=>{const last=i===all.length-1;d.querySelector('[data-a=r]').hidden=!last||!!ctrl;if(!last)d.querySelector('[data-a=k]').hidden=true})}
async function regen(){if(ctrl||!chat.length||chat[chat.length-1].role!=='assistant')return;chat.pop();render();await run()}
log.onscroll=()=>{$('jump').hidden=stick()};$('jump').onclick=()=>{log.scrollTop=log.scrollHeight};
document.addEventListener('keydown',e=>{if(e.key==='Escape'&&ctrl)ctrl.abort();if((e.ctrlKey||e.metaKey)&&e.key.toLowerCase()==='k'){e.preventDefault();newChat()}});
function busy(b){go.textContent=b?'Stop':'Send';go.classList.toggle('stop',b)}
async function send(){
 let text=ta.value.trim();if(!text&&!att)return;
 hush();const u={role:'user',content:text,shown:text};
 if(att&&att.img){u.images=[att.img];u.content=text||'Describe this picture in detail.';u.shown=`${text}\n[Picture: ${att.name}]`.trim();att=null;$('hint').textContent='Enter to send, Shift+Enter for a new line'}else if(att){u.content=`File ${att.name}:\n${att.text}\n\n${text}`;u.shown=`${text}\n[Attached: ${att.name}]`.trim();att=null;$('hint').textContent='Enter to send, Shift+Enter for a new line'}
 if(!chat.length)log.innerHTML='';
 chat.push(u);ta.value='';ta.style.height='auto';addMsg(u);save();await run()}
let raf=0,lastSave=0;
function stick(){return log.scrollHeight-log.scrollTop-log.clientHeight<80}
function queueDraw(d,a){if(raf)return;raf=requestAnimationFrame(()=>{raf=0;const s=stick();paint(d,a,true);if(s)log.scrollTop=log.scrollHeight})}
async function run(){
 const a={role:'assistant',content:'',think:''};chat.push(a);const d=addMsg(a,true);busy(true);
 ctrl=new AbortController();const t0=Date.now();
 try{const r=await fetch('/api/chat',{method:'POST',signal:ctrl.signal,headers:{'Content-Type':'application/json'},
   body:JSON.stringify({model:sel.value,effort:thk.value,ctx:+$('ctx').value,tools:$('tools').checked,messages:chat.slice(0,-1).map(({role,content,images},i,a)=>i===a.length-1?{role,content,images}:{role,content})})});
  const rd=r.body.getReader(),dec=new TextDecoder();let buf='';
  while(true){const{value,done}=await rd.read();if(done)break;buf+=dec.decode(value,{stream:true});const parts=buf.split('\n\n');buf=parts.pop();
   for(const p of parts){if(!p.startsWith('data: '))continue;const e=JSON.parse(p.slice(6));
    if(e.t)a.think+=e.t;if(e.c){if(!a.content)a.secs=Math.round((Date.now()-t0)/1000);a.content+=e.c}
    if(e.stats)a.stats=e.stats;if(e.err)a.err=e.err;queueDraw(d,a);feed(a,false);if(Date.now()-lastSave>1500){lastSave=Date.now();save()}}}
 }catch(x){if(x.name==='AbortError')a.stats='Stopped';else a.err='Connection lost: '+x.message}
 cancelAnimationFrame(raf);raf=0;paint(d,a,false);save();busy(false);ctrl=null;tidy();ta.focus();if(a.stats!=='Stopped')feed(a,true)}
const HINT='Enter to send, Shift+Enter for a new line',dk=$('dark');
dk.checked=document.documentElement.dataset.theme==='dark';
dk.onchange=()=>{const t=dk.checked?'dark':'light';document.documentElement.dataset.theme=t;localStorage.theme=t};
async function copyText(t){try{await navigator.clipboard.writeText(t);return true}catch(_){const x=document.createElement('textarea');x.value=t;x.style.cssText='position:fixed;opacity:0';document.body.appendChild(x);x.select();let ok=false;try{ok=document.execCommand('copy')}catch(_){}x.remove();return ok}}
function note(s){$('hint').textContent=s;setTimeout(()=>{if(!att)$('hint').textContent=HINT},4000)}
async function dl(kind,msgs){const r=await fetch('/api/export',{method:'POST',headers:{'Content-Type':'application/json'},body:JSON.stringify({format:kind,messages:msgs})});
 if(!r.ok)return note(await r.text());
 const name=(/filename="([^"]+)"/.exec(r.headers.get('Content-Disposition')||'')||[])[1]||'export';
 const url=URL.createObjectURL(await r.blob()),a=document.createElement('a');a.href=url;a.download=name;document.body.appendChild(a);a.click();a.remove();setTimeout(()=>URL.revokeObjectURL(url),4000)}
log.onclick=async e=>{const b=e.target.closest('.cp');if(!b)return;if(await copyText(b.closest('.cb').querySelector('code').textContent)){b.textContent='Copied';setTimeout(()=>{b.textContent='Copy'},1200)}};
document.querySelectorAll('[data-x]').forEach(b=>b.onclick=()=>{const m=chat.filter(x=>x.content).map(x=>({role:x.role,content:x.shown??x.content}));if(!m.length)return note('Nothing to export yet');dl(b.dataset.x,m)});
let au=null,queue=[],pumping=false,gen=0;
function hush(){gen++;queue=[];pumping=false;if(au){au.pause();au=null}}
function getAudio(t){return fetch('/api/tts',{method:'POST',headers:{'Content-Type':'application/json'},body:JSON.stringify({text:t,voice:vsel.value,speed:+vspd.value})})
 .then(r=>r.ok?r.blob():null).then(b=>b&&b.size?URL.createObjectURL(b):null).catch(()=>null)}
function enqueue(t){queue.push(t);if(!pumping)pump()}
async function pump(){pumping=true;const g=gen;let next=queue.length?getAudio(queue.shift()):null;
 while(next){const url=await next;if(g!==gen)return;next=queue.length?getAudio(queue.shift()):null;
  if(url)await new Promise(done=>{au=new Audio(url);au.onended=au.onerror=done;au.play().catch(done)});if(g!==gen)return}
 pumping=false;if(queue.length)pump()}
function speak(text){hush();text.split(/\n\s*\n/).filter(p=>p.trim()).forEach(enqueue)}
function feed(a,final){if(!spkChk.checked||!a.content)return;const from=a.sp||0,rest=a.content.slice(from);let cut=0;
 if(final)cut=rest.length;else{const all=[...rest.matchAll(/[.!?]["')\]]?\s|\n\n/g)];if(all.length){const m=all[all.length-1];cut=m.index+m[0].length}}
 if(!cut||(!final&&cut<40))return;
 if(!final&&((a.content.slice(0,from+cut).match(/```/g)||[]).length)%2)return;
 a.sp=from+cut;const piece=rest.slice(0,cut);if(piece.trim())enqueue(piece)}
go.onclick=()=>ctrl?ctrl.abort():send();
ta.onkeydown=e=>{if(e.key==='Enter'&&!e.shiftKey){e.preventDefault();if(!ctrl)send()}
 if(e.key==='ArrowUp'&&!ta.value&&!ctrl){const i=chat.map(m=>m.role).lastIndexOf('user');if(i>=0){e.preventDefault();ta.value=chat[i].shown??chat[i].content;chat.length=i;save();render();ta.oninput()}}};
ta.oninput=()=>{ta.style.height='auto';ta.style.height=Math.min(ta.scrollHeight,200)+'px'};
$('new').onclick=newChat;
$('attach').onclick=()=>$('file').click();
const vis=$('vis');vis.checked=localStorage.vis==='1';vis.onchange=()=>localStorage.vis=vis.checked?'1':'0';
function shrink(f){return new Promise((ok,no)=>{const im=new Image(),u=URL.createObjectURL(f);im.onload=()=>{const k=Math.min(1,768/Math.max(im.width,im.height)),c=document.createElement('canvas');c.width=Math.round(im.width*k);c.height=Math.round(im.height*k);c.getContext('2d').drawImage(im,0,0,c.width,c.height);URL.revokeObjectURL(u);ok(c.toDataURL('image/jpeg',0.85).split(',')[1])};im.onerror=()=>no(new Error('Could not open the picture'));im.src=u})}
const tl=$('tools');tl.checked=localStorage.tools==='1';tl.onchange=()=>localStorage.tools=tl.checked?'1':'0';
$('ctx').value=localStorage.ctx||'16384';$('ctx').onchange=()=>localStorage.ctx=$('ctx').value;
$('file').onchange=async e=>{const f=e.target.files[0];e.target.value='';if(!f)return;const h=$('hint');let text;
 try{if(f.type.startsWith('image/')&&vis.checked){h.textContent='Preparing picture...';att={name:f.name,img:await shrink(f)};h.textContent='Attached (the model will look at it): '+f.name;return}
 if(f.type.startsWith('image/')){h.textContent='Reading text from the picture...';
   const r=await fetch('/api/ocr',{method:'POST',headers:{'Content-Type':f.type},body:f});
   if(!r.ok){h.textContent=await r.text();return}text=await r.text();if(!text.trim()){h.textContent='No text found in that picture';return}}
  else{if(f.size>2000000){h.textContent='File too large (2 MB max)';return}text=await f.text()}}
 catch(x){h.textContent='Could not read the file: '+x.message;return}
 const cap=Math.floor(+$('ctx').value*0.6*3.5);let note='';if(text.length>cap){text=text.slice(0,cap);note=' (cut to fit; raise Memory to read more)'}
 att={name:f.name,text};h.textContent='Attached: '+f.name+note};
fetch('/api/models').then(r=>r.json()).then(j=>{sel.innerHTML=j.models.map(n=>`<option>${n}</option>`).join('');
 const want=localStorage.model||j.default;if(j.models.includes(want))sel.value=want});
function setTitle(){const c=chats.find(x=>x.id===cur);$('ttl').textContent=c?c.title:'Local assistant'}
$('menu').onclick=()=>document.body.classList.toggle('menu');$('scrim').onclick=()=>document.body.classList.remove('menu');$('nw').onclick=newChat;
let tx=0,ty=0;addEventListener('touchstart',e=>{tx=e.touches[0].clientX;ty=e.touches[0].clientY},{passive:true});
addEventListener('touchend',e=>{const dx=e.changedTouches[0].clientX-tx,dy=Math.abs(e.changedTouches[0].clientY-ty),b=document.body.classList;
 if(dy>50)return;if(tx<28&&dx>70)b.add('menu');else if(dx<-70&&b.contains('menu'))b.remove('menu')},{passive:true});
const vsel=$('vsel'),vspd=$('vspd');vspd.value=localStorage.vspd||1;vspd.onchange=()=>localStorage.vspd=vspd.value;
fetch('/api/voices').then(r=>r.json()).then(j=>{vsel.innerHTML=j.voices.map(n=>`<option value="${n}">${n.replace('.onnx','').replace('kokoro:','Kokoro ')}</option>`).join('');
 const want=localStorage.voice||j.default;if(j.voices.includes(want))vsel.value=want;vsel.onchange=()=>localStorage.voice=vsel.value});
sel.onchange=()=>localStorage.model=sel.value;render();
</script></body></html>'''


class Handler(BaseHTTPRequestHandler):
    def log_message(self, *args):
        pass

    def send(self, code, ctype, body, headers=()):
        self.send_response(code)
        self.send_header("Content-Type", ctype)
        self.send_header("Content-Length", str(len(body)))
        for k, v in headers:
            self.send_header(k, v)
        self.end_headers()
        self.wfile.write(body)

    def ok(self):
        if not TOKEN:
            return True
        cookies = [c.strip() for c in self.headers.get("Cookie", "").split(";")]
        return any(secrets.compare_digest(c, f"ct={TOKEN}") for c in cookies)

    def do_GET(self):
        u = urlparse(self.path)
        deny = lambda: self.send(401, "text/plain", b"Open this page once with ?token=YOUR_TOKEN")
        if u.path == "/":
            if TOKEN and secrets.compare_digest(parse_qs(u.query).get("token", [""])[0], TOKEN):
                cookie = [("Set-Cookie", f"ct={TOKEN}; Path=/; HttpOnly; SameSite=Strict")]
            elif self.ok():
                cookie = []
            else:
                return deny()
            return self.send(200, "text/html; charset=utf-8", PAGE.encode(), cookie)
        if not self.ok():
            return deny()
        if u.path == "/api/models":
            try:
                names = [m["name"] for m in requests.get(OLLAMA + "/api/tags", timeout=5).json()["models"]]
            except Exception:
                names = []
            body = json.dumps({"models": names, "default": CONFIG.get("model", "")}).encode()
            return self.send(200, "application/json", body)
        if u.path == "/api/voices":
            return self.send(200, "application/json", json.dumps({"voices": kokoro_names() + voice_files(), "default": kokoro_default() or os.path.basename(VOICE)}).encode())
        self.send(404, "text/plain", b"Not found")

    def event(self, obj):
        self.wfile.write(f"data: {json.dumps(obj)}\n\n".encode())
        self.wfile.flush()

    def tts(self):
        try:
            body = json.loads(self.rfile.read(int(self.headers.get("Content-Length", "0"))))
            text = speech_text(body.get("text", ""))
            voice = os.path.join(BASE, body["voice"]) if body.get("voice") in voice_files() else VOICE
            speed = min(1.6, max(0.6, float(body.get("speed") or 1)))
            cmd = [sys.executable, "-m", "piper", "--model", voice] + ([] if speed == 1 else ["--length-scale", f"{1 / speed:.2f}"])
            if not text:
                return self.send(204, "audio/wav", b"")
            if str(body.get("voice", "")) in kokoro_names():
                voice = body["voice"]
            key = hashlib.sha1(f"{voice}|{speed}|{text}".encode()).hexdigest()
            hit = os.path.join(CACHE, key + ".wav")
            if os.path.exists(hit):
                with open(hit, "rb") as f:
                    return self.send(200, "audio/wav", f.read())
            if voice.startswith("kokoro:"):
                data = kokoro_wav(text, voice[7:], speed)
                save_cache(hit, data)
                return self.send(200, "audio/wav", data)
            with tempfile.NamedTemporaryFile(suffix=".wav") as wav:
                subprocess.run(cmd + ["--output_file", wav.name],
                               input=text.encode("utf-8"), check=True, capture_output=True, timeout=180)
                with open(wav.name, "rb") as f:
                    data = f.read()
                save_cache(hit, data)
                self.send(200, "audio/wav", data)
        except Exception as error:
            self.send(500, "text/plain", str(error).encode())

    def ocr(self):
        try:
            size = int(self.headers.get("Content-Length", "0"))
            if size > 15 * 1024 * 1024:
                return self.send(413, "text/plain", b"Image too large (15 MB max)")
            data = self.rfile.read(size)
            with tempfile.NamedTemporaryFile() as img:
                img.write(data)
                img.flush()
                out = subprocess.run(["tesseract", img.name, "stdout", "-l", os.environ.get("OCR_LANG", "eng")],
                                     capture_output=True, timeout=120)
            if out.returncode != 0:
                return self.send(500, "text/plain", out.stderr[:300] or b"OCR failed")
            self.send(200, "text/plain; charset=utf-8", out.stdout.decode("utf-8", errors="replace").strip().encode())
        except FileNotFoundError:
            self.send(501, "text/plain", b"OCR needs Tesseract. Run: sudo apt install tesseract-ocr")
        except Exception as error:
            self.send(500, "text/plain", f"OCR failed: {error}".encode())

    def export(self):
        try:
            size = int(self.headers.get("Content-Length", "0"))
            if size > MAX_EXPORT:
                return self.send(413, "text/plain", b"Chat is too large to export")
            req = json.loads(self.rfile.read(size))
            fmt, msgs = req.get("format"), clean_messages(req.get("messages"))
            if not msgs:
                return self.send(400, "text/plain", b"Nothing to export")
            stamp, title = datetime.now().strftime("%Y%m%d-%H%M"), f"Chat {datetime.now():%Y-%m-%d %H:%M}"
            if fmt == "md":
                data, name, ctype = to_markdown(msgs, title).encode("utf-8"), f"chat-{stamp}.md", "text/markdown; charset=utf-8"
            elif fmt == "docx":
                data, name = build_docx(msgs, title), f"chat-{stamp}.docx"
                ctype = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            elif fmt in ("code", "bundle"):
                files = code_blocks(msgs)
                if fmt == "code" and not files:
                    return self.send(404, "text/plain", b"No code blocks found in this chat")
                entries = [(("code/" if fmt == "bundle" else "") + n, t) for n, t in files]
                if fmt == "bundle":
                    entries.insert(0, ("chat.md", to_markdown(msgs, title)))
                    try:
                        entries.insert(1, ("chat.docx", build_docx(msgs, title)))
                    except ImportError:
                        pass  # Word export is optional
                data, name, ctype = make_zip(entries), f"{fmt}-{stamp}.zip", "application/zip"
            else:
                return self.send(400, "text/plain", b"Unknown format")
            self.send(200, ctype, data, [("Content-Disposition", f'attachment; filename="{name}"')])
        except ImportError:
            self.send(501, "text/plain", b"Word export needs python-docx. Run: pip install python-docx")
        except Exception as error:
            self.send(500, "text/plain", f"Export failed: {error}".encode())

    def do_POST(self):
        path = urlparse(self.path).path
        if path not in ("/api/chat", "/api/tts", "/api/export", "/api/ocr") or not self.ok():
            return self.send(404, "text/plain", b"Not found")
        if path == "/api/tts":
            return self.tts()
        if path == "/api/export":
            return self.export()
        if path == "/api/ocr":
            return self.ocr()
        upstream = None
        try:
            req = json.loads(self.rfile.read(int(self.headers.get("Content-Length", "0"))))
            sm = system_message()
            msgs = ([sm] if sm else []) + [m for m in req.get("messages", [])[-60:] if m.get("role") in ("user", "assistant")]
            ctx = min(max(int(req.get("ctx") or CONFIG.get("num_ctx", 16384)), 2048), 65536)
            msgs = fit(msgs, ctx)
            payload = {"model": req.get("model") or CONFIG.get("model"), "messages": msgs, "stream": True,
                       "think": decide_think(req, msgs, req.get("model") or CONFIG.get("model")), "keep_alive": "30m", "options": {**CONFIG.get("generation", {}), **gpu_options(req.get("model") or CONFIG.get("model")), "num_ctx": ctx}}
            if any(m.get("images") for m in msgs) and not can_see(payload["model"]):
                self.send_response(200)
                self.send_header("Content-Type", "text/event-stream")
                self.send_header("Cache-Control", "no-cache")
                self.end_headers()
                return self.event({"err": payload["model"] + " cannot look at pictures. Pick the 9B or heretic-writer, or turn off 'Model sees pictures' to use text-reading (OCR) instead."})
            stats = ""
            headers_sent = False
            for _round in range(4):  # the model may call tools a few times before answering
                if req.get("tools"):
                    payload["tools"] = TOOLS
                upstream = requests.post(OLLAMA + "/api/chat", json=payload, stream=True, timeout=(10, 900))
                if not headers_sent:
                    self.send_response(200)
                    self.send_header("Content-Type", "text/event-stream")
                    self.send_header("Cache-Control", "no-cache")
                    self.end_headers()
                    headers_sent = True
                if upstream.status_code != 200:
                    return self.event({"err": f"Ollama said: {upstream.text[:300]}"})
                calls, text = [], ""
                for line in upstream.iter_lines():
                    if not line:
                        continue
                    chunk = json.loads(line)
                    m = chunk.get("message", {})
                    if m.get("thinking"):
                        self.event({"t": m["thinking"]})
                    if m.get("content"):
                        text += m["content"]
                        self.event({"c": m["content"]})
                    calls += m.get("tool_calls") or []
                    if chunk.get("done"):
                        n, ns = chunk.get("eval_count", 0), chunk.get("eval_duration", 0)
                        stats = f"{n / (ns / 1e9):.1f} tokens/s" if ns else ""
                        break
                upstream.close()
                if not calls:
                    break
                payload["messages"].append({"role": "assistant", "content": text, "tool_calls": calls})
                for call in calls:
                    func = call.get("function", {})
                    name, args = func.get("name", ""), func.get("arguments") or {}
                    self.event({"t": f"\nUsing {name}({json.dumps(args)[:120]})\n"})
                    payload["messages"].append({"role": "tool", "tool_name": name, "content": tool_reply(name, args)})
            self.event({"stats": stats})
        except (BrokenPipeError, ConnectionResetError):
            pass  # browser pressed Stop or closed the tab
        except Exception as error:
            try:
                self.event({"err": f"Error: {error}"})
            except OSError:
                pass
        finally:
            if upstream is not None:
                upstream.close()


if __name__ == "__main__":
    import socket
    try:  # find the LAN address (no traffic is sent)
        with socket.socket(socket.AF_INET, socket.SOCK_DGRAM) as probe:
            probe.connect(("10.255.255.255", 1))
            lan = probe.getsockname()[0]
    except OSError:
        lan = "<laptop-ip>"
    suffix = f"/?token={TOKEN}" if TOKEN else ""
    print(f"Chat UI on this machine: http://127.0.0.1:{PORT}{suffix}")
    if HOST not in ("127.0.0.1", "localhost"):
        print(f"Chat UI on your network: http://{lan}:{PORT}{suffix}")
        if not TOKEN:
            print("WARNING: no password set. Anyone on this network can chat and export. Set CHAT_TOKEN=yourpassword to lock it.")
    ThreadingHTTPServer((HOST, PORT), Handler).serve_forever()
